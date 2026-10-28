from __future__ import annotations

import json
import re
import shutil
import subprocess
import tempfile
from copy import deepcopy
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Pt
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_ROOT = ROOT / "data" / "templates"
OUTPUT_ROOT = ROOT / "var" / "artifacts"
MAX_PAGES = 8
MAX_BRIEF_CHARS = 500


class DocumentError(RuntimeError):
    def __init__(self, message: str, *, removable_question_ids: list[str] | None = None):
        super().__init__(message)
        self.removable_question_ids = removable_question_ids or []


def template_manifest() -> dict[str, Any]:
    path = TEMPLATE_ROOT / "manifest.json"
    if not path.exists():
        raise DocumentError("公司模板尚未匯入")
    return json.loads(path.read_text(encoding="utf-8"))


def _soffice() -> str:
    candidates = [
        shutil.which("soffice.com"), shutil.which("soffice"),
        r"C:\Program Files\LibreOffice\program\soffice.com",
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]
    for value in candidates:
        if value and Path(value).exists():
            return str(value)
    raise DocumentError("找不到 LibreOffice，無法驗證八頁限制")


def _apply_style(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "標楷體"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.15


def _append_case_content(document: Document, case: dict[str, Any], role: str) -> None:
    heading = document.add_paragraph()
    heading_run = heading.add_run("本案補充問答")
    heading_run.bold = True
    heading_run.font.size = Pt(16)
    intake = case["intake"]
    document.add_paragraph(f"發生（破獲）時間：{intake['event_time']}")
    document.add_paragraph(f"發生（破獲）地點：{intake['location']}")
    document.add_paragraph(f"案情摘要：{intake['summary']}")
    answers = case.get("answers", {}).get(role, {})
    question_lookup = {
        item["question_id"]: item for batch in case.get("question_batches", []) for item in batch.get("questions", [])
    }
    for item in case.get("next_question_batch", []):
        question_lookup.setdefault(item["question_id"], item)
    dynamic_answers = [(qid, answer) for qid, answer in answers.items() if not qid.startswith("T-")]
    for index, (question_id, answer) in enumerate(dynamic_answers, start=1):
        question = question_lookup.get(question_id, {"text": question_id})["text"]
        document.add_paragraph(f"問{index}：{question}")
        document.add_paragraph(f"答：{answer}")
    document.add_paragraph("以上筆錄經受詢問人閱覽或朗讀確認，如有增刪更正應於簽名或按指印前完成。")


def _fill_company_answers(document: Document, role: str, answers: dict[str, str]) -> None:
    questions_path = TEMPLATE_ROOT / "questions.json"
    if not questions_path.exists():
        raise DocumentError("公司模板拆題 binding 不存在")
    bindings = json.loads(questions_path.read_text(encoding="utf-8"))["questions"]
    grouped: dict[tuple, list[tuple[int, str]]] = {}
    original_by_binding: dict[tuple, str] = {}
    for question in bindings:
        if question["role"] != role or question["question_id"] not in answers:
            continue
        binding = question["template_binding"]
        key = tuple((name, binding[name]) for name in sorted(binding))
        grouped.setdefault(key, []).append((question["atomic_index"], answers[question["question_id"]]))
        original_by_binding[key] = question["source_original_text"]
    for key, atomic_answers in grouped.items():
        binding = dict(key)
        answer_text = "；".join(text for _, text in sorted(atomic_answers))
        if binding.get("kind") == "table_cell":
            cell = document.tables[int(binding["table"])].rows[int(binding["row"])].cells[int(binding["cell"])]
            if "答：" not in cell.text:
                cell.add_paragraph(f"答：{answer_text}")
        elif binding.get("kind") == "paragraph":
            original = original_by_binding[key]
            paragraph = next((item for item in document.paragraphs if item.text.strip() == original), None)
            if paragraph is not None:
                paragraph.add_run(f"\n答：{answer_text}")


def _page_count(docx_path: Path) -> int:
    with tempfile.TemporaryDirectory(prefix="cia-render-") as temp_value:
        temp = Path(temp_value)
        profile = temp / "profile"
        command = [
            _soffice(), "--headless", f"-env:UserInstallation={profile.as_uri()}",
            "--convert-to", "pdf", "--outdir", str(temp), str(docx_path),
        ]
        try:
            result = subprocess.run(command, capture_output=True, text=True, timeout=90, check=False)
        except subprocess.TimeoutExpired as exc:
            raise DocumentError("LibreOffice 頁數驗證逾時") from exc
        pdf = temp / f"{docx_path.stem}.pdf"
        if result.returncode != 0 or not pdf.exists():
            raise DocumentError("LibreOffice 無法完成頁數驗證")
        return len(PdfReader(str(pdf)).pages)


def build_interview(case: dict[str, Any], role: str) -> dict[str, Any]:
    template_role = "transfer" if role == "transfer" else role
    manifest = template_manifest()
    entry = manifest["templates"].get(template_role)
    if entry is None:
        raise DocumentError(f"沒有 {template_role} 模板")
    template = TEMPLATE_ROOT / entry["converted_path"]
    document = Document(template)
    _apply_style(document)
    answer_role = "suspect" if role == "transfer" else role
    _fill_company_answers(document, template_role, case.get("answers", {}).get(answer_role, {}))
    _append_case_content(document, case, answer_role)
    OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)
    path = OUTPUT_ROOT / f"{case['case_id']}-r{case['revision']}-{role}.docx"
    document.save(path)
    pages = _page_count(path)
    if pages > MAX_PAGES:
        optional = [
            item["question_id"] for batch in case.get("question_batches", [])
            for item in batch.get("questions", []) if item.get("requirement") != "REQUIRED"
        ]
        path.unlink(missing_ok=True)
        raise DocumentError(
            f"筆錄實際排版為 {pages} 頁，超過 {MAX_PAGES} 頁，已禁止下載。",
            removable_question_ids=optional,
        )
    return {
        "path": str(path), "pages": pages, "template_role": template_role,
        "template_sha256": entry["converted_sha256"], "revision": case["revision"],
    }


def roc_text(value: str) -> str:
    def replace(match: re.Match[str]) -> str:
        year = int(match.group(1))
        return f"民國{year - 1911}年"
    return re.sub(r"(?<!民國)(20\d{2})年", replace, value)


def brief_text(case: dict[str, Any]) -> str:
    intake = case["intake"]
    candidates = [item for item in case["candidates"] if item["status"] != "UNSUPPORTED_ROUTE"]
    reason = candidates[0]["label"] if candidates else "待查證案件"
    numbers = case.get("case_numbers", {})
    number_parts = []
    labels = {"proof_number": "證明單案號", "women_child_number": "婦幼案號", "case_165_number": "165案號"}
    for key in ("proof_number", "women_child_number", "case_165_number"):
        if numbers.get(key):
            number_parts.append(f"{labels[key]}：{numbers[key]}")
    number_text = "、".join(number_parts)
    lines = [
        f"案由：{reason}",
        f"發生（破獲）時間：{roc_text(intake['event_time'])}",
        f"發生（破獲）地點：{intake['location']}",
        f"依規定陳報（{number_text}）" if number_text else "依規定陳報",
        f"案情摘要：{roc_text(intake['summary'])}",
    ]
    result = "\n".join(lines)
    if len(result) > MAX_BRIEF_CHARS:
        raise DocumentError(f"條陳共 {len(result)} 字，超過 {MAX_BRIEF_CHARS} 字，請重新摘要。")
    return result


def build_brief(case: dict[str, Any]) -> dict[str, Any]:
    text = brief_text(case)
    document = Document()
    _apply_style(document)
    document.add_heading("條陳", level=1)
    for line in text.splitlines():
        document.add_paragraph(line)
    OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)
    path = OUTPUT_ROOT / f"{case['case_id']}-r{case['revision']}-brief.docx"
    document.save(path)
    return {"path": str(path), "text": text, "chars": len(text), "revision": case["revision"]}
